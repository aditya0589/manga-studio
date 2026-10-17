import os
import uuid
import io
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from pypdf import PdfReader
from docx import Document as DocxDocument
from PIL import Image

from app.core.config import settings
from app.core.database import get_db
from app.models.user import User
from app.models.project import Project
from app.models.story import StoryDocument
from app.models.asset import Asset
from app.schemas.story import StoryDocumentResponse
from app.api.dependencies import get_current_user, get_project_with_ownership
from app.storage import storage, compute_checksum

router = APIRouter(tags=["Uploads"])

def extract_text_from_file(content: bytes, filename: str, mime_type: str) -> str:
    """Extract clean text from TXT, PDF, or DOCX files."""
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    try:
        if ext == ".txt" or "text/plain" in mime_type:
            text = content.decode("utf-8", errors="replace")
        elif ext == ".pdf" or "pdf" in mime_type:
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for p in reader.pages:
                t = p.extract_text()
                if t:
                    pages_text.append(t)
            text = "\n".join(pages_text)
        elif ext in (".docx", ".doc") or "word" in mime_type:
            doc = DocxDocument(io.BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs if p.text])
        else:
            # Fallback text decoding
            text = content.decode("utf-8", errors="ignore")
    except Exception as e:
        text = f"Error extracting document text: {e}"
    return text.strip()

@router.post("/projects/{project_id}/uploads", response_model=StoryDocumentResponse)
async def upload_document(
    project_id: str,
    file: UploadFile = File(...),
    document_type: str = Form("manuscript"), # manuscript, character_notes, world_lore
    project: Project = Depends(get_project_with_ownership),
    db: AsyncSession = Depends(get_db)
):
    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)
    if file_size_mb > settings.MAX_UPLOAD_SIZE_MB:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum allowed upload size of {settings.MAX_UPLOAD_SIZE_MB}MB"
        )

    # Save to storage
    ext = os.path.splitext(file.filename or "file")[1]
    safe_filename = f"{uuid.uuid4().hex}{ext}"
    storage_key = f"projects/{project_id}/documents/{safe_filename}"
    saved_key = await storage.save(content, storage_key, file.content_type or "application/octet-stream")

    # Extract text if document
    extracted_text = extract_text_from_file(content, file.filename or "", file.content_type or "")

    doc_record = StoryDocument(
        project_id=project_id,
        original_filename=file.filename or "uploaded_file",
        mime_type=file.content_type or "application/octet-stream",
        storage_key=saved_key,
        extracted_text=extracted_text,
        processing_status="processed" if extracted_text else "uploaded"
    )
    db.add(doc_record)
    await db.commit()
    await db.refresh(doc_record)
    return StoryDocumentResponse.model_validate(doc_record)

@router.get("/projects/{project_id}/documents", response_model=List[StoryDocumentResponse])
async def list_project_documents(
    project: Project = Depends(get_project_with_ownership),
    db: AsyncSession = Depends(get_db)
):
    stmt = select(StoryDocument).where(StoryDocument.project_id == project.id).order_by(StoryDocument.created_at)
    res = await db.execute(stmt)
    docs = res.scalars().all()
    return [StoryDocumentResponse.model_validate(d) for d in docs]

@router.get("/documents/{document_id}", response_model=StoryDocumentResponse)
async def get_document(
    document_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    doc = await db.get(StoryDocument, document_id)
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    project = await db.get(Project, doc.project_id)
    if not project or project.user_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")

    return StoryDocumentResponse.model_validate(doc)

@router.delete("/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    document_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    doc = await db.get(StoryDocument, document_id)
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    project = await db.get(Project, doc.project_id)
    if not project or project.user_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")

    await storage.delete(doc.storage_key)
    await db.delete(doc)
    await db.commit()
    return None
